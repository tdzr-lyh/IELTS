from __future__ import annotations

import hashlib
import json
import re
import shutil
from pathlib import Path

import pdfplumber
from docx import Document


PROJECT_ROOT = Path(__file__).resolve().parent
MATERIAL_ROOT = Path(r"D:\tdzr\myself\雅思\线上\听力\听力材料")
OUTPUT_DATA = PROJECT_ROOT / "listening-data.js"
OUTPUT_MEDIA = PROJECT_ROOT / "media" / "listening"

SCENE_ROOT = MATERIAL_ROOT / "听力场景高频单词"
JIJING_ROOT = MATERIAL_ROOT / "听力机经词汇"

SCENE_PDFS = [
    "1. 租房场景机经高频词汇.pdf",
    "2. 旅游场景机经高频词汇.pdf",
    "3. 工作求职场景.pdf",
    "4. 休闲活动场景机经高频词汇.pdf",
    "5. 地图题词汇汇总.pdf",
    "6 银行场景.pdf",
    "7 图书馆场景.pdf",
    "8 博物馆场景.pdf",
    "9 餐饮场景.pdf",
    "10 医疗场景.pdf",
    "11. Part 2 旅游类题干词.pdf",
    "12 Part 2 城市规划高频题干词.pdf",
    "13. Part 2 工作机构类题干词.pdf",
    "14. Part2 生活服务类高频题干词.pdf",
    "15. Part 2 休闲场所活动题干词.pdf",
    "16. 学术场景机经高频词汇.pdf",
    "17. 动植物机经高频词汇.pdf",
    "18. 企业类高频单词.pdf",
    "19. 环境场景高频单词.pdf",
    "20. 文化场景高频单词.pdf",
]

JIJING_FILES = [
    Path("Chapter1 数字浪涛/Chapter 1 数字淘沙.docx"),
    Path("Chapter2 生存必杀/Chapter2 生存必杀/Chapter 2 生存必杀.docx"),
    Path("Chapter3 事无巨细/Chapter3 事无巨细/Chapter 3 事无巨细.docx"),
    Path("Chapter4 地北天南/Chapter4 地北天南/Chapter 4 地北天南.docx"),
    Path("Chapter5 学术不羁/Chapter 5 学术不羁.docx"),
    Path("Chapter6 动物世界/Chapter6 动物世界/Chapter 6 动物世界.docx"),
    Path("Chapter7 历史追溯/Chapter7 历史追溯.docx"),
    Path("Chapter8 心有灵犀/Chapter8 心有灵犀.docx"),
    Path("Chapter9 商管并进/Chapter9 商管并进/Chapter 9 商管并进.docx"),
    Path("Chapter10 科学达人/Chapter 10 科学达人.docx"),
    Path("Chapter11 新题补丁/Chapter 11 新题补丁.docx"),
    Path("Chapter 12 新题补丁/Chapter 12 新题补丁.docx"),
    Path("Chapter13新题补丁/Chapter 13 新题补丁.docx"),
    Path("Chapter14 查漏补缺/chapter14.pdf"),
    Path("Chapter15 稳固推进/Chapter 15.pdf"),
]

TITLE_REPLACEMENTS = {
    "1. 租房场景机经高频词汇": "租房与住宿",
    "2. 旅游场景机经高频词汇": "旅游与出行",
    "3. 工作求职场景": "工作与求职",
    "4. 休闲活动场景机经高频词汇": "休闲与活动",
    "5. 地图题词汇汇总": "地图与方位",
    "6 银行场景": "银行与金融",
    "7 图书馆场景": "图书馆",
    "8 博物馆场景": "博物馆",
    "9 餐饮场景": "餐饮",
    "10 医疗场景": "医疗与健康",
    "11. Part 2 旅游类题干词": "Part 2 旅游题干",
    "12 Part 2 城市规划高频题干词": "Part 2 城市规划",
    "13. Part 2 工作机构类题干词": "Part 2 工作机构",
    "14. Part2 生活服务类高频题干词": "Part 2 生活服务",
    "15. Part 2 休闲场所活动题干词": "Part 2 休闲场所",
    "16. 学术场景机经高频词汇": "学术场景",
    "17. 动植物机经高频词汇": "动植物",
    "18. 企业类高频单词": "企业与管理",
    "19. 环境场景高频单词": "环境",
    "20. 文化场景高频单词": "文化",
}

POS_PATTERN = re.compile(
    r"\b(?:adj|adv|n|v|prep|conj|pron|num|phr|vt|vi)\s*\.?\b",
    re.IGNORECASE,
)
PHONETIC_PATTERN = re.compile(r"/[^/\n]{1,80}/")
LEADING_PATTERN = re.compile(
    r"^\s*(?:[（(]?\d{1,3}[）).、．]?\s*|[【\[]拓[】\]]\s*|[-•·]+\s*)"
)
CHINESE_PATTERN = re.compile(r"[\u3400-\u9fff]")
LATIN_PATTERN = re.compile(r"[A-Za-z]")
PAGE_ONLY_PATTERN = re.compile(r"^\s*\d{1,3}\s*$")


def slug(value: str) -> str:
    value = value.lower().strip()
    value = re.sub(r"[^a-z0-9]+", "-", value).strip("-")
    return value or hashlib.sha1(value.encode("utf-8")).hexdigest()[:10]


def normalize_spaces(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def split_english_chinese(raw: str) -> tuple[str, str] | None:
    line = normalize_spaces(raw.replace("\u3000", " "))
    if not line or PAGE_ONLY_PATTERN.match(line):
        return None
    if re.fullmatch(r"[［\[][^\]］]{1,45}[\]］][:：]?", line):
        return None
    line = LEADING_PATTERN.sub("", line)
    line = PHONETIC_PATTERN.sub(" ", line)
    line = POS_PATTERN.sub(" ", line)
    line = normalize_spaces(line)
    if not LATIN_PATTERN.search(line):
        return None

    chinese_match = CHINESE_PATTERN.search(line)
    if chinese_match:
        english = line[: chinese_match.start()].strip(" ：:，,；;。.[]【】()（）")
        meaning = line[chinese_match.start() :].strip()
    else:
        english = line.strip(" ：:，,；;。.[]【】()（）")
        meaning = ""

    english = re.sub(r"^[A-Z][A-Za-z ]{0,35}[:：]\s*", "", english)
    english = (
        english.replace("（", "(")
        .replace("）", ")")
        .replace("／", "/")
        .replace("…", " ")
    )
    english = re.sub(r"[(){}\[\]［］]", " ", english)
    english = normalize_spaces(english)
    if not english or not LATIN_PATTERN.search(english):
        return None
    if len(english) > 72:
        return None
    lowered = english.lower()
    banned = (
        "approach",
        "雅思",
        "section one",
        "section 1",
        "chapter",
        "level",
        "copyright",
        "sky123",
        "环球天下",
    )
    if any(token.lower() in lowered for token in banned):
        return None
    if lowered in {
        "person",
        "area",
        "location",
        "facilities",
        "furniture",
        "diet",
        "fees",
        "part2",
        "part 2",
        "第一组",
        "第二组",
    }:
        return None
    return english, meaning


def make_entry(term: str, meaning: str, source: str, group: str) -> dict[str, str]:
    stable = f"{group}|{source}|{term.lower()}"
    return {
        "id": hashlib.sha1(stable.encode("utf-8")).hexdigest()[:14],
        "term": term,
        "meaning": meaning,
        "source": source,
        "group": group,
    }


def dedupe(entries: list[dict[str, str]]) -> list[dict[str, str]]:
    result: list[dict[str, str]] = []
    seen: set[str] = set()
    for item in entries:
        key = re.sub(r"[^a-z0-9]+", "", item["term"].lower())
        if len(key) < 2 or key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


def extract_pdf_entries(path: Path, title: str, group: str) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for raw in text.splitlines():
                parsed = split_english_chinese(raw)
                if not parsed:
                    continue
                term, meaning = parsed
                entries.append(make_entry(term, meaning, title, group))
    return dedupe(entries)


def docx_chunks(path: Path) -> list[str]:
    document = Document(path)
    chunks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return chunks


def extract_docx_entries(path: Path, title: str) -> list[dict[str, str]]:
    entries: list[dict[str, str]] = []
    for chunk in docx_chunks(path):
        for raw in chunk.splitlines():
            parsed = split_english_chinese(raw)
            if not parsed:
                continue
            term, meaning = parsed
            entries.append(make_entry(term, meaning, title, "jijing"))
    return dedupe(entries)


def category_title(path: Path) -> str:
    stem = path.stem
    if stem in TITLE_REPLACEMENTS:
        return TITLE_REPLACEMENTS[stem]
    chapter = re.search(r"chapter\s*(\d+)", stem, re.IGNORECASE)
    if chapter:
        number = int(chapter.group(1))
        parent = path.parent.name
        parent = re.sub(r"^Chapter\s*\d+\s*", "", parent, flags=re.IGNORECASE).strip()
        return f"Chapter {number} · {parent or '机经词汇'}"
    return stem


def build_categories() -> tuple[list[dict], list[dict]]:
    scene_categories = []
    for filename in SCENE_PDFS:
        path = SCENE_ROOT / filename
        title = category_title(path)
        entries = extract_pdf_entries(path, title, "scene")
        if entries:
            scene_categories.append(
                {
                    "id": f"scene-{len(scene_categories) + 1:02d}",
                    "title": title,
                    "sourceFile": filename,
                    "entries": entries,
                }
            )

    jijing_categories = []
    for relative in JIJING_FILES:
        path = JIJING_ROOT / relative
        title = category_title(path)
        if path.suffix.lower() == ".docx":
            entries = extract_docx_entries(path, title)
        else:
            entries = extract_pdf_entries(path, title, "jijing")
        if entries:
            jijing_categories.append(
                {
                    "id": f"jijing-{len(jijing_categories) + 1:02d}",
                    "title": title,
                    "sourceFile": str(relative).replace("\\", "/"),
                    "entries": entries,
                }
            )
    return scene_categories, jijing_categories


def classify_track(relative: Path) -> tuple[str, str]:
    text = str(relative)
    level = "5.5" if "5.5分听写" in text else "6.0"
    track_type = "basic" if "听力基本功训练" in text else "vocabulary"
    return level, track_type


def track_title(path: Path) -> str:
    pieces = list(path.parts)
    filename = path.stem
    parent = pieces[-2] if len(pieces) >= 2 else ""
    if parent.endswith("音频") or parent.startswith("听力"):
        return filename
    return f"{parent} · {filename}"


def build_tracks() -> list[dict]:
    OUTPUT_MEDIA.mkdir(parents=True, exist_ok=True)
    for old_file in OUTPUT_MEDIA.glob("*"):
        if old_file.is_file():
            old_file.unlink()

    tracks = []
    roots = [
        MATERIAL_ROOT / "5.5分听写",
        MATERIAL_ROOT / "6.0分听写",
    ]
    counters: dict[tuple[str, str], int] = {}
    for root in roots:
        for source in sorted(root.rglob("*.mp3")):
            relative = source.relative_to(MATERIAL_ROOT)
            level, track_type = classify_track(relative)
            key = (level, track_type)
            counters[key] = counters.get(key, 0) + 1
            safe_level = level.replace(".", "")
            filename = f"{safe_level}-{track_type}-{counters[key]:02d}.mp3"
            target = OUTPUT_MEDIA / filename
            shutil.copy2(source, target)
            tracks.append(
                {
                    "id": filename.removesuffix(".mp3"),
                    "level": level,
                    "type": track_type,
                    "title": track_title(relative),
                    "sourceFile": str(relative).replace("\\", "/"),
                    "url": f"./media/listening/{filename}",
                }
            )
    return tracks


def write_data(scene_categories: list[dict], jijing_categories: list[dict], tracks: list[dict]) -> None:
    payload = {
        "generatedFrom": str(MATERIAL_ROOT),
        "sceneCategories": scene_categories,
        "jijingCategories": jijing_categories,
        "dictationTracks": tracks,
    }
    json_text = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    OUTPUT_DATA.write_text(
        "// 由“听力材料”文件夹自动生成，请运行 build_listening_materials.py 更新。\n"
        f"window.LISTENING_LIBRARY = {json_text};\n",
        encoding="utf-8",
    )


def main() -> None:
    if not MATERIAL_ROOT.exists():
        raise FileNotFoundError(MATERIAL_ROOT)
    scene_categories, jijing_categories = build_categories()
    tracks = build_tracks()
    write_data(scene_categories, jijing_categories, tracks)
    print(f"Scene categories: {len(scene_categories)}")
    print(f"Scene entries: {sum(len(item['entries']) for item in scene_categories)}")
    print(f"Jijing categories: {len(jijing_categories)}")
    print(f"Jijing entries: {sum(len(item['entries']) for item in jijing_categories)}")
    print(f"Dictation tracks: {len(tracks)}")
    print(f"Data: {OUTPUT_DATA}")
    print(f"Media: {OUTPUT_MEDIA}")


if __name__ == "__main__":
    main()
